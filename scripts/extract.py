"""Ekstrak data SWS W36 - G2 ke JSON untuk dashboard statis.
Kolom NO TELP PIC dibuang; nomor telepon di teks catatan di-redaksi."""
import openpyxl, json, re, os
from datetime import datetime, date
from collections import Counter, defaultdict

SRC = r"C:\Users\ASUS\.kimi-code\sessions\wd_kimi_82d34ab51235\session_873a9658-3bfb-4455-998e-ffe4958e6082\attachments\f_70e4423c-bfa3-4f57-a710-5d6d328bf032-SWS_W36_-_G2_MASTER_UPDATE_7_AREA_FINAL (1).xlsx"
OUT = os.path.join(os.path.dirname(__file__), "..", "data")
os.makedirs(OUT, exist_ok=True)

wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)

def dt(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return None

def num(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return v
    return None

def s(v):
    if v is None:
        return None
    t = str(v).strip()
    return t if t else None

# ---------------- EXT ----------------
ext_rows = []
for r in wb["EXT"].iter_rows(min_row=5, values_only=True):
    if not r[7]:
        continue
    omset25 = [num(r[c]) or 0 for c in range(60, 72)]
    omset26 = [num(r[c]) or 0 for c in range(73, 85)]
    ext_rows.append({
        "no": num(r[1]), "grfpm": s(r[2]), "region": s(r[3]), "kota": s(r[4]),
        "kecamatan": s(r[5]), "afps": s(r[6]), "outlet": s(r[7]),
        "jenis": s(r[9]), "alamat": s(r[11]), "pic": s(r[12]),
        "channel": s(r[14]), "kode_outlet": s(r[15]),
        "takeover": s(r[17]), "kompetitor": s(r[18]), "bbbrbl": s(r[19]),
        "brand": s(r[20]), "nomor_ap": s(r[21]),
        "start": dt(r[22]), "end": dt(r[23]),
        "kompensasi": num(r[24]), "produksi": num(r[25]), "total_nilai": num(r[26]),
        "plan_perpanjang": s(r[33]), "hasil_visit": s(r[38]),
        "tgl_deal": dt(r[39]), "week_deal": num(r[40]), "reason_not_ext": s(r[41]),
        "omset_week": num(r[42]),
        "nomor_ap_new": s(r[46]), "start_new": dt(r[47]), "end_new": dt(r[48]),
        "kompensasi_new": num(r[49]),
        "omset25": omset25, "omset26": omset26,
        "total25": num(r[72]), "total26": num(r[85]),
    })
print("EXT:", len(ext_rows))

# ---------------- NOO ----------------
noo_rows = []
for r in wb["NOO"].iter_rows(min_row=4, values_only=True):
    if not r[6]:
        continue
    noo_rows.append({
        "grfpm": s(r[1]), "region": s(r[2]), "kota": s(r[3]), "kecamatan": s(r[4]),
        "afps": s(r[5]), "outlet": s(r[6]), "jenis": s(r[8]),
        "prioritas": s(r[9]), "kategori_kpi": s(r[10]),
        "alamat": s(r[12]), "pic": s(r[13]),  # no telp (col O) sengaja dibuang
        "channel": s(r[15]), "rating": num(r[16]), "perating": num(r[17]),
        "score": num(r[18]), "siswa": num(r[19]), "omset_week": num(r[20]),
        "alasan": s(r[21]), "historis": s(r[22]),
        "week_visit": num(r[23]), "hari_visit": s(r[24]),
    })
print("NOO:", len(noo_rows))

# ---------------- CATATAN UPDATE ----------------
PHONE_RE = re.compile(r"(\+?62[\s\-]?|0)8\d{2}[\s\-]?\d{4}[\s\-]?\d{2,5}")
def redact(t):
    return PHONE_RE.sub("[redacted]", t)

catatan = []
for name in ["CATATAN UPDATE", "CATATAN UPDATE 2", "CATATAN UPDATE 3", "CATATAN UPDATE 4"]:
    ws = wb[name]
    blocks = []
    cur = []
    for r in ws.iter_rows(values_only=True):
        cells = [redact(str(v).strip()) for v in r if v is not None and str(v).strip()]
        line = " | ".join(cells)
        if line:
            cur.append(line)
        elif cur:
            blocks.append("\n".join(cur)); cur = []
    if cur:
        blocks.append("\n".join(cur))
    catatan.append({"sheet": name, "blocks": blocks})
print("CATATAN sheets:", len(catatan))

# ---------------- SUMMARY ----------------
def agg(rows, key, val_keys):
    d = defaultdict(lambda: {k: 0 for k in val_keys} | {"count": 0})
    for r in rows:
        k = r.get(key) or "-"
        d[k]["count"] += 1
        for vk in val_keys:
            v = r.get(vk)
            if isinstance(v, (int, float)):
                d[k][vk] += v
    return {k: v for k, v in sorted(d.items(), key=lambda x: -x[1]["count"])}

today = date(2026, 9, 23)
def days_until(iso):
    if not iso:
        return None
    try:
        return (datetime.strptime(iso, "%Y-%m-%d").date() - today).days
    except ValueError:
        return None

visit_status = Counter()
for r in ext_rows:
    h = (r["hasil_visit"] or "").strip().upper()
    if not h:
        visit_status["BELUM DIVISIT"] += 1
    elif "DEAL" in h and "NO" not in h:
        visit_status["DEAL"] += 1
    elif "NO DEAL" in h:
        visit_status["NO DEAL"] += 1
    elif "PROSES" in h:
        visit_status["PROSES"] += 1
    else:
        visit_status[h] += 1

exp_buckets = {"expired": 0, "d30": 0, "d60": 0, "d90": 0, "d180": 0, "later": 0, "unknown": 0}
for r in ext_rows:
    d = days_until(r["end_new"] or r["end"])
    if d is None: exp_buckets["unknown"] += 1
    elif d < 0: exp_buckets["expired"] += 1
    elif d <= 30: exp_buckets["d30"] += 1
    elif d <= 60: exp_buckets["d60"] += 1
    elif d <= 90: exp_buckets["d90"] += 1
    elif d <= 180: exp_buckets["d180"] += 1
    else: exp_buckets["later"] += 1

noo_historis = Counter()
for r in noo_rows:
    h = (r["historis"] or "").strip().upper()
    if not h: noo_historis["BELUM ADA"] += 1
    elif "DEAL" in h: noo_historis["ADA PROGRESS/DEAL"] += 1
    else: noo_historis["PERNAH DILIBATKAN"] += 1

monthly = []
for i in range(12):
    monthly.append({
        "bulan": f"2025-{i+1:02d}",
        "omset": sum(r["omset25"][i] for r in ext_rows),
    })
for i in range(12):
    monthly.append({
        "bulan": f"2026-{i+1:02d}",
        "omset": sum(r["omset26"][i] for r in ext_rows),
    })

summary = {
    "generated": today.isoformat(),
    "kpi": {
        "total_ext": len(ext_rows),
        "total_noo": len(noo_rows),
        "total_kompensasi": sum(r["kompensasi"] or 0 for r in ext_rows),
        "total_omset_2025": sum(r["total25"] or 0 for r in ext_rows),
        "total_omset_2026": sum(r["total26"] or 0 for r in ext_rows),
        "takeover": sum(1 for r in ext_rows if r["takeover"] and "BUKAN" not in r["takeover"].upper()),
        "prioritas_noo": sum(1 for r in noo_rows if r["prioritas"] == "PRIORITAS"),
    },
    "per_region": agg(ext_rows, "region", ["kompensasi", "total25", "total26"]),
    "per_kota": agg(ext_rows, "kota", ["kompensasi"]),
    "per_jenis": agg(ext_rows, "jenis", ["kompensasi"]),
    "per_brand": agg(ext_rows, "brand", ["kompensasi"]),
    "noo_per_region": agg(noo_rows, "region", []),
    "noo_per_kategori": agg(noo_rows, "kategori_kpi", []),
    "visit_status": dict(visit_status.most_common()),
    "expiring": exp_buckets,
    "noo_historis": dict(noo_historis.most_common()),
    "monthly_omset": monthly,
}

with open(os.path.join(OUT, "ext.json"), "w", encoding="utf-8") as f:
    json.dump(ext_rows, f, ensure_ascii=False, separators=(",", ":"))
with open(os.path.join(OUT, "noo.json"), "w", encoding="utf-8") as f:
    json.dump(noo_rows, f, ensure_ascii=False, separators=(",", ":"))
with open(os.path.join(OUT, "catatan.json"), "w", encoding="utf-8") as f:
    json.dump(catatan, f, ensure_ascii=False, separators=(",", ":"))
with open(os.path.join(OUT, "summary.json"), "w", encoding="utf-8") as f:
    json.dump(summary, f, ensure_ascii=False, indent=1)

print("kpi:", json.dumps(summary["kpi"]))
print("visit_status:", dict(visit_status.most_common()))
print("expiring:", exp_buckets)
print("noo_historis:", dict(noo_historis.most_common()))
for fn in os.listdir(OUT):
    print(fn, os.path.getsize(os.path.join(OUT, fn)))
